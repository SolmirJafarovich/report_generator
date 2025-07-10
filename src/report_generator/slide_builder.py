from pptx import Presentation

from .image_slide import add_image_slide
from .table_slide import add_table_slide
from .text_slide import add_text_slide
from .title_slide import add_title_slide
from .utils.text_config import TextConfig


class SlideBuilder:
    def __init__(self, presentation: Presentation, text_config: TextConfig):
        self.prs = presentation
        self.text_config = text_config

    def add_slide(self, slide_data: dict):
        slide_type = slide_data.get("type", "text")

        if slide_type == "title":
            add_title_slide(self, slide_data)
        elif slide_type == "image":
            add_image_slide(self, slide_data)
        elif slide_type == "table":
            add_table_slide(self, slide_data)
        else:
            add_text_slide(self, slide_data)
